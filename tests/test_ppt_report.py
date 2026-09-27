import asyncio
import json
import struct
import zlib
from copy import deepcopy
from io import BytesIO

import pytest
from pydantic import ValidationError

pytest.importorskip("pptx")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.reporting.ai import ReportAiService
from app.reporting.context import build_report_context
from app.reporting.ppt_ai import (
    PPT_DIAGNOSIS_SYSTEM_PROMPT,
    PPT_PROMPT_VERSION,
    PPT_REPORT_KIND,
    PptDiagnosis,
    PptRecommendedAction,
    generate_ppt_diagnosis,
)
from app.reporting.ppt_report import build_diagnostic_ppt, diagnostic_ppt_filename_stem


def tiny_png() -> bytes:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + kind
            + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\x1c\x91\xa8")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")


def compact_context() -> dict:
    return {
        "reportKind": PPT_REPORT_KIND,
        "contextVersion": "cms-report-context-v1",
        "evidence": {
            "equipment": {
                "id": "V032",
                "displayName": "No.2 Direct Cooling Pump / Pump DS",
                "plant": "Water Treatment",
                "measurementPoint": "Pump DS",
                "sensorNumber": "32",
                "rpm": 1780,
                "bearing": "6318",
                "status": "caution",
                "lastMeasuredAt": "2026-01-15 12:00:00",
            },
            "trend": {
                "available": True,
                "analysis": {"status": "rising", "slope": 0.01},
                "metrics": [
                    {
                        "key": "velocity",
                        "label": "속도",
                        "latestValue": 2.0,
                        "unit": "mm/s",
                        "warningThreshold": 1.5,
                        "dangerThreshold": 3.0,
                    },
                    {
                        "key": "accelerationAll",
                        "label": "가속도 ALL",
                        "latestValue": 0.3,
                        "unit": "g",
                        "warningThreshold": 0.5,
                        "dangerThreshold": 1.0,
                    },
                ],
            },
            "fft": {
                "available": True,
                "measuredAt": "2026-01-15 12:00:00",
                "measurement": {"measureMode": {"label": "속도"}},
                "diagnosis": {"label": "Unbalance", "status": "analyzed"},
                "spectrum": {"peaks": [{"frequency": 30.0, "amplitude": 1.2}]},
            },
        },
    }


def diagnosis() -> PptDiagnosis:
    return PptDiagnosis(
        headline="Trend 상승과 1X 성분을 추가 확인해야 합니다.",
        conclusion="속도 값과 FFT Peak를 현장 운전 조건과 함께 재확인하는 것이 필요합니다.",
        trend_finding="최근 Trend는 상승 상태이며 속도 값이 주의 기준에 근접했습니다.",
        fft_finding="30 Hz 부근의 1X Peak가 관측되어 회전 성분을 추가 확인해야 합니다.",
        recommended_actions=[
            PptRecommendedAction(priority="high", action="체결 상태와 운전 조건을 현장 점검합니다."),
            PptRecommendedAction(priority="planned", action="동일 조건에서 FFT를 재측정합니다."),
        ],
        limitations=["운전 부하 정보가 없어 원인 해석에는 한계가 있습니다."],
    )


def visible_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def test_build_diagnostic_ppt_fills_the_reference_template_with_chart_images() -> None:
    output = build_diagnostic_ppt(compact_context(), diagnosis(), tiny_png(), tiny_png())
    presentation = Presentation(BytesIO(output))

    assert len(presentation.slides) == 1
    assert presentation.slide_width == 12192000
    assert presentation.slide_height == 6858000
    slide = presentation.slides[0]
    text = visible_text(slide)
    tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
    metadata = tables[1]
    assert "[Water Treatment] CMS Data 모니터링 및 진단결과" in text
    assert metadata.cell(1, 1).text == "No.2 Direct Cooling Pump / Pump DS"
    assert metadata.cell(0, 1).text == "Water Treatment"
    assert metadata.cell(0, 8).text == "’26.1.15 12:00"
    assert metadata.cell(1, 11).text == "32"
    assert metadata.cell(2, 2).text == "1,780 RPM"
    assert metadata.cell(2, 9).text == "1.5 / 3"
    assert metadata.cell(2, 12).text == "0.5 / 1"
    assert metadata.cell(3, 1).text.startswith("속도 2 mm/s, 가속도 ALL 0.3 g.")
    assert metadata.cell(4, 1).text == diagnosis().headline
    assert "체결 상태와 운전 조건을 현장 점검합니다." in text
    assert "Bearing 내륜 결함에 의한 충격" not in text
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes) >= 2
    assert diagnostic_ppt_filename_stem(compact_context()) == "cms-diagnostic-V032-20260115"


def test_build_diagnostic_ppt_leaves_missing_source_data_blank() -> None:
    context = deepcopy(compact_context())
    del context["evidence"]["equipment"]["displayName"]
    del context["evidence"]["equipment"]["plant"]
    del context["evidence"]["equipment"]["sensorNumber"]
    del context["evidence"]["fft"]["measuredAt"]
    context["evidence"]["trend"]["metrics"] = []
    output = build_diagnostic_ppt(context, diagnosis(), None, tiny_png())
    presentation = Presentation(BytesIO(output))

    slide = presentation.slides[0]
    tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
    metadata = tables[1]
    assert "정보 없음" not in visible_text(slide)
    assert metadata.cell(0, 1).text == ""
    assert metadata.cell(0, 8).text == ""
    assert metadata.cell(1, 1).text == ""
    assert metadata.cell(1, 11).text == ""
    assert metadata.cell(2, 9).text == ""
    assert metadata.cell(2, 12).text == ""
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes) == 1


def test_ppt_diagnosis_rejects_html_text() -> None:
    payload = diagnosis().model_dump()
    payload["headline"] = "<strong>근거 없는 강조</strong>"

    with pytest.raises(ValidationError):
        PptDiagnosis.model_validate(payload)


class CapturingGateway:
    def __init__(self) -> None:
        self.calls: list[tuple[str, str, dict, type[PptDiagnosis]]] = []

    async def generate(self, report_kind, instructions, context, response_model):
        self.calls.append((report_kind, instructions, context, response_model))
        return response_model.model_validate(diagnosis().model_dump())


def test_ppt_ai_uses_distinct_kind_and_serializes_no_raw_spectrum_arrays() -> None:
    context = build_report_context(
        PPT_REPORT_KIND,
        "1",
        {"equipment": {"id": "V032"}},
        {
            "recordId": 7,
            "frequencies": [10.0, 20.0],
            "amplitudes": [0.1, 0.2],
            "analysis": {},
        },
        [],
    )
    context["evidence"]["fft"]["frequencies"] = [10.0, 20.0]
    context["evidence"]["fft"]["amplitudes"] = [0.1, 0.2]

    gateway = CapturingGateway()
    result = asyncio.run(
        generate_ppt_diagnosis(ReportAiService(gateway, "test-model"), context)
    )

    report_kind, instructions, prompt_context, response_model = gateway.calls[0]
    serialized = json.dumps(prompt_context)
    assert result.headline == diagnosis().headline
    assert report_kind == PPT_REPORT_KIND
    assert report_kind != "email-draft"
    assert response_model is PptDiagnosis
    assert PPT_PROMPT_VERSION != "cms-email-report-v1"
    assert "frequencies" not in serialized
    assert "amplitudes" not in serialized
    assert "RUL" in instructions
    assert "MTBF" in instructions
    assert "HTML" in PPT_DIAGNOSIS_SYSTEM_PROMPT
