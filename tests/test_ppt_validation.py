import asyncio
from copy import deepcopy
from io import BytesIO
from urllib.parse import unquote
import json

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError

from app.ai_analysis import AiProviderError, OpenAISettings
from app.main import create_app
from app.reporting.ai import ReportAiService
from app.reporting.ppt_layout import PptLayout, default_ppt_layout
from app.reporting.ppt_pipeline import PptPipeline, PptRepair, PptTextEdit, apply_repair
from app.reporting.ppt_report import _fit_png_to_frame, build_diagnostic_ppt
from app.reporting.ppt_validation import (
    OpenAIPptGateway, PptIssue, PptRenderError, PptReview, geometry_issues, inspect_ppt,
)
from app.reporting import ppt_validation
from tests.test_app import make_data_dir, post
from tests.test_ppt_report import compact_context, diagnosis, tiny_png
from tests.test_report_endpoints import FakeReportGateway, png_data_url


def issue(description="설명이 잘렸습니다."):
    return PptIssue(element_id="trend_text", kind="clipping", severity="error",
                    description=description, suggestion="설명 영역의 높이를 확보하십시오.")


def changed_layout(height=0.95):
    payload = default_ppt_layout().model_dump()
    payload["trend_text"]["height"] = height
    return PptLayout.model_validate(payload)


class ScriptedGateway:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    async def generate(self, kind, instructions, context, model):
        self.calls.append((kind, context))
        value = self.responses.pop(0)
        if isinstance(value, Exception):
            raise value
        assert isinstance(value, model)
        return value


class Renderer:
    def __init__(self, fail_at=None):
        self.presentations = []
        self.fail_at = fail_at

    async def __call__(self, presentation):
        self.presentations.append(presentation)
        if len(self.presentations) == self.fail_at:
            raise PptRenderError("renderer failed")
        return tiny_png()


def run_pipeline(responses, renderer=None):
    gateway = ScriptedGateway(responses)
    renderer = renderer or Renderer()
    pipeline = PptPipeline(ReportAiService(gateway, "test-model"), renderer)
    result = asyncio.run(pipeline.build(compact_context(), diagnosis(), tiny_png(), tiny_png()))
    return result, gateway, renderer


def test_initial_slide_has_no_geometry_errors_and_merged_cells_are_inspected():
    manifest = inspect_ppt(build_diagnostic_ppt(compact_context(), diagnosis(), tiny_png(), tiny_png()))
    assert geometry_issues(manifest) == []
    metadata = next(item for item in manifest["elements"] if item["element_id"] == "metadata")
    equipment = next(cell for cell in metadata["cells"] if cell["element_id"] == "metadata:1:1")
    assert equipment["text"] == compact_context()["evidence"]["equipment"]["displayName"]
    assert equipment["width"] > 4


def test_geometry_checks_boundary_and_nonintentional_overlap():
    deck = Presentation(BytesIO(build_diagnostic_ppt(compact_context(), diagnosis(), tiny_png(), tiny_png())))
    slide = deck.slides[0]
    next(shape for shape in slide.shapes if shape.name == "title").left = Inches(-0.1)
    trend = next(shape for shape in slide.shapes if shape.name == "trend_text")
    fft = next(shape for shape in slide.shapes if shape.name == "fft_text")
    fft.left, fft.top = trend.left, trend.top
    output = BytesIO()
    deck.save(output)
    issues = geometry_issues(inspect_ppt(output.getvalue()))
    assert any(item.element_id == "title" and item.kind == "boundary" for item in issues)
    assert any(item.element_id == "trend_text" and item.kind == "overlap" for item in issues)


@pytest.mark.parametrize("field,value", [("x", -1), ("width", 50), ("font_size", 9), ("height", float("nan"))])
def test_layout_rejects_unsafe_repairs(field, value):
    payload = default_ppt_layout().model_dump()
    payload["trend_text"][field] = value
    with pytest.raises(ValidationError):
        PptLayout.model_validate(payload)


@pytest.mark.parametrize("size", [(1000, 100), (100, 1000), (800, 400)])
def test_chart_fit_preserves_all_edges_and_aspect_ratio(size):
    source = Image.new("RGB", size, "red")
    draw = ImageDraw.Draw(source)
    corner_width, corner_height = max(5, size[0] // 10), max(5, size[1] // 10)
    for x, y, color in [(0, 0, "lime"), (size[0] - corner_width, 0, "blue"),
                        (0, size[1] - corner_height, "yellow"),
                        (size[0] - corner_width, size[1] - corner_height, "black")]:
        draw.rectangle((x, y, x + corner_width - 1, y + corner_height - 1), fill=color)
    stream = BytesIO()
    source.save(stream, format="PNG")
    with Image.open(_fit_png_to_frame(stream.getvalue(), 585, 250)) as fitted:
        colors = fitted.getcolors(fitted.width * fitted.height)
        assert (0, 255, 0) in [color for count, color in colors]
        assert (255, 0, 0) in [color for count, color in colors]
        assert all(color in [rgb for count, rgb in colors] for color in [(0, 0, 255), (255, 255, 0), (0, 0, 0)])
        assert fitted.width / fitted.height == pytest.approx(585 / 250, rel=0.003)


def test_direct_pass_uses_real_image_content_and_only_one_review():
    result, gateway, renderer = run_pipeline([PptReview(issues=[])])
    assert result.validation.status == "passed"
    assert result.validation.repairs == 0
    assert len(renderer.presentations) == len(gateway.calls) == 1
    context = gateway.calls[0][1]
    assert context["pptHash"]
    content = OpenAIPptGateway(OpenAISettings(None, "test", 1))._user_content(context)
    assert content[1]["type"] == "image_url"
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")
    assert "slideImage" not in content[0]["text"]


def test_repair_then_pass_revalidates_changed_artifact():
    result, gateway, renderer = run_pipeline([
        PptReview(issues=[issue()]), PptRepair(layout=changed_layout(), text_edits=[]), PptReview(issues=[]),
    ])
    assert result.validation.status == "passed"
    assert result.validation.repairs == 1
    assert [kind for kind, _ in gateway.calls] == ["ppt-ui-validation", "ppt-ui-repair", "ppt-ui-validation"]
    assert gateway.calls[0][1]["pptHash"] != gateway.calls[2][1]["pptHash"]
    assert result.presentation == renderer.presentations[-1]


def test_two_repairs_limit_and_best_candidate_selection():
    result, gateway, renderer = run_pipeline([
        PptReview(issues=[issue()]), PptRepair(layout=changed_layout(), text_edits=[]),
        PptReview(issues=[issue(), issue("두 번째 문제") .model_copy(update={"kind": "readability"})]),
        PptRepair(layout=changed_layout(1.0), text_edits=[]),
        PptReview(issues=[issue(), issue().model_copy(update={"kind": "chart"})]),
    ])
    assert result.validation.status == "warning"
    assert result.validation.repairs == 2
    assert len(renderer.presentations) == 3
    assert len(gateway.calls) == 5
    assert result.presentation == renderer.presentations[0]


def test_no_change_stops_without_second_review():
    result, gateway, renderer = run_pipeline([
        PptReview(issues=[issue()]), PptRepair(layout=default_ppt_layout(), text_edits=[]),
    ])
    assert result.validation.status == "warning"
    assert len(renderer.presentations) == 1
    assert len(gateway.calls) == 2


def test_unknown_element_is_unverified_not_passed():
    result, _, _ = run_pipeline([PptReview(issues=[issue().model_copy(update={"element_id": "made-up"})])])
    assert result.validation.status == "unverified"


@pytest.mark.parametrize("failure", [AiProviderError(), PptRenderError("unavailable")])
def test_validation_failure_keeps_downloadable_ppt(failure):
    renderer = Renderer(fail_at=1) if isinstance(failure, PptRenderError) else Renderer()
    result, _, _ = run_pipeline([failure], renderer)
    assert result.validation.status == "unverified"
    assert len(Presentation(BytesIO(result.presentation)).slides) == 1


def test_later_renderer_failure_preserves_previously_reviewed_candidate():
    result, _, renderer = run_pipeline([
        PptReview(issues=[issue()]), PptRepair(layout=changed_layout(), text_edits=[]),
    ], Renderer(fail_at=2))
    assert result.validation.status == "unverified"
    assert result.presentation == renderer.presentations[0]


def test_invalid_text_repair_preserves_previous_artifact():
    result, _, renderer = run_pipeline([
        PptReview(issues=[issue()]), PptRepair(layout=changed_layout(), text_edits=[
            PptTextEdit(field="trend_finding", text="속도는 999 mm/s입니다."),
        ]),
    ])
    assert result.validation.status == "warning"
    assert result.presentation == renderer.presentations[0]
    assert len(renderer.presentations) == 1


def test_invalid_layout_response_preserves_reviewed_artifact():
    result, _, renderer = run_pipeline([PptReview(issues=[issue()]), ValueError("invalid layout")])
    assert result.validation.status == "warning"
    assert result.presentation == renderer.presentations[0]


def test_unexpected_geometry_blocks_visual_agent_pass():
    renderer = Renderer()
    gateway = ScriptedGateway([
        PptReview(issues=[issue()]),
        PptRepair(layout=default_ppt_layout().model_copy(update={
            "fft_text": default_ppt_layout().trend_text,
        }), text_edits=[]),
        PptReview(issues=[]), PptRepair(layout=changed_layout(), text_edits=[]), PptReview(issues=[]),
    ])
    result = asyncio.run(PptPipeline(ReportAiService(gateway, "test"), renderer).build(
        compact_context(), diagnosis(), tiny_png(), tiny_png()))
    assert result.validation.status == "passed"
    assert result.validation.repairs == 2
    assert len(renderer.presentations) == 3


def test_same_artifact_reuses_visual_review_but_force_bypasses_cache():
    async def run():
        gateway = ScriptedGateway([PptReview(issues=[]), PptReview(issues=[])])
        pipeline = PptPipeline(ReportAiService(gateway, "test"), Renderer())
        await pipeline.build(compact_context(), diagnosis(), tiny_png(), tiny_png())
        await pipeline.build(compact_context(), diagnosis(), tiny_png(), tiny_png())
        await pipeline.build(compact_context(), diagnosis(), tiny_png(), tiny_png(), force=True)
        return gateway.calls
    assert len(asyncio.run(run())) == 2


@pytest.mark.parametrize("fail", [False, True])
def test_renderer_isolates_profiles_and_cleans_temp_files(monkeypatch, fail):
    directories = []
    profiles = []
    monkeypatch.setattr(ppt_validation.shutil, "which", lambda name: "/tools/" + name)

    async def fake_command(*arguments, cwd):
        directories.append(cwd)
        if "libreoffice" in arguments[0]:
            profiles.append(arguments[1])
            (cwd / "slide.pdf").write_bytes(b"pdf")
        elif fail:
            raise PptRenderError("render failure")
        else:
            (cwd / "slide.png").write_bytes(tiny_png())

    monkeypatch.setattr(ppt_validation, "_run_renderer", fake_command)

    async def run():
        return await asyncio.gather(ppt_validation.render_ppt(b"ppt"), ppt_validation.render_ppt(b"ppt"), return_exceptions=True)

    results = asyncio.run(run())
    assert len(set(profiles)) == 2
    assert all(not directory.exists() for directory in directories)
    if fail:
        assert all(isinstance(result, PptRenderError) for result in results)
    else:
        assert results == [tiny_png(), tiny_png()]


def test_renderer_missing_tools_is_explicit(monkeypatch):
    monkeypatch.setattr(ppt_validation.shutil, "which", lambda name: None)
    with pytest.raises(PptRenderError, match="설치"):
        asyncio.run(ppt_validation.render_ppt(b"ppt"))


def test_renderer_timeout_kills_process_group(monkeypatch, tmp_path):
    killed = []

    class Process:
        pid = 99999
        returncode = None
        waits = 0

        async def wait(self):
            self.waits += 1
            if self.waits == 1:
                raise TimeoutError()
            self.returncode = -9

    async def spawn(*arguments, **kwargs):
        assert kwargs["start_new_session"] is True
        return Process()

    monkeypatch.setattr(ppt_validation.asyncio, "create_subprocess_exec", spawn)
    monkeypatch.setattr(ppt_validation.os, "killpg", lambda pid, sig: killed.append(pid))
    with pytest.raises(PptRenderError):
        asyncio.run(ppt_validation._run_renderer("test", cwd=tmp_path))
    assert killed == [99999]


def test_extract_only_summary_preserves_numbers_and_uncertainty():
    original = diagnosis().model_copy(update={"trend_finding": "속도는 2 mm/s입니다. 운전 정보가 없어 해석에 한계가 있습니다. 운전 정보가 없어 해석에 한계가 있습니다."})
    shortened = "속도는 2 mm/s입니다. 운전 정보가 없어 해석에 한계가 있습니다."
    repair = PptRepair(layout=default_ppt_layout(), text_edits=[PptTextEdit(field="trend_finding", text=shortened)])
    updated = apply_repair(original, repair)
    assert updated.trend_finding == shortened
    assert updated.recommended_actions == original.recommended_actions
    with pytest.raises(ValueError):
        apply_repair(original, PptRepair(layout=default_ppt_layout(), text_edits=[
            PptTextEdit(field="trend_finding", text="속도는 2 mm/s입니다."),
        ]))


def test_repair_cannot_drop_non_numeric_korean_evidence():
    original = diagnosis().model_copy(update={"trend_finding": "진동이 상승했습니다. 베어링 결함 성분이 관측됩니다."})
    with pytest.raises(ValueError, match="진단 근거"):
        apply_repair(original, PptRepair(layout=default_ppt_layout(), text_edits=[
            PptTextEdit(field="trend_finding", text="진동이 상승했습니다."),
        ]))


def test_long_text_is_not_silently_truncated():
    context = deepcopy(compact_context())
    context["evidence"]["equipment"]["displayName"] = "긴 설비 이름 " * 25
    long = diagnosis().model_copy(update={"trend_finding": "관측 결과를 동일 조건에서 추가 확인합니다. " * 10})
    deck = Presentation(BytesIO(build_diagnostic_ppt(context, long, tiny_png(), tiny_png())))
    slide = deck.slides[0]
    assert next(shape for shape in slide.shapes if shape.name == "trend_text").text == long.trend_finding
    metadata = next(shape for shape in slide.shapes if shape.name == "metadata")
    assert metadata.table.cell(1, 1).text == context["evidence"]["equipment"]["displayName"].strip()


@pytest.mark.parametrize("status", ["passed", "warning", "unverified"])
def test_endpoint_returns_download_and_validation_metadata(tmp_path, status):
    response_list = [PptReview(issues=[])] if status == "passed" else [
        PptReview(issues=[issue()]), PptRepair(layout=default_ppt_layout(), text_edits=[]),
    ]
    renderer = Renderer(fail_at=1 if status == "unverified" else None)
    response = post(create_app(make_data_dir(tmp_path), report_gateway=FakeReportGateway(),
                               ppt_gateway=ScriptedGateway(response_list), ppt_renderer=renderer),
                    "/api/dashboard/1/diagnostic-ppt", {
                        "recordId": 2, "trendImageDataUrl": png_data_url(), "fftImageDataUrl": png_data_url(),
                    })
    assert response.status_code == 200
    metadata = json.loads(unquote(response.headers["X-PPT-Validation"]))
    assert metadata["status"] == status
    assert len(Presentation(BytesIO(response.content)).slides) == 1
