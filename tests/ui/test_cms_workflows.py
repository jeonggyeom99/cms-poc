from __future__ import annotations

from io import BytesIO

import pytest

from tests.ui.helpers import capture_page, open_page


pytestmark = pytest.mark.ui


def test_local_report_download_and_email(cms_page, cms_side_effects, tmp_path, monkeypatch):
    from playwright.sync_api import expect
    from pptx import Presentation
    import dotenv

    def deny_dotenv(*args, **kwargs):
        raise AssertionError("UI fixture schema validation must not load .env")

    monkeypatch.setattr(dotenv, "load_dotenv", deny_dotenv)
    from app.ai_analysis import AiAnalysisResponse
    from app.reporting.email_draft import EmailDraftResponse

    open_page(cms_page, "/dashboard2?dataset=1&until=2026-10-29")
    export = cms_page.locator("#diagnostic-ppt-export")
    email = cms_page.locator("#email-draft-open")
    expect(export).to_be_enabled()
    expect(email).to_be_enabled()
    ai = AiAnalysisResponse.model_validate(cms_side_effects.responses["ai-analysis"])
    real = cms_page.request.get("/api/dashboard/1", params={"until": "2026-10-29"}).json()
    assert ai.analysis.fft_reference_record_ids == [real["fft"]["recordId"]]
    assert real["fft"]["recordId"] in [record["id"] for record in real["fft"]["records"]]
    with cms_page.expect_download() as downloaded:
        export.click()
    download = downloaded.value
    assert download.suggested_filename == "diagnostic-ppt-1.pptx"
    target = tmp_path / download.suggested_filename
    download.save_as(target)
    presentation = Presentation(BytesIO(target.read_bytes()))
    assert len(presentation.slides) == 1
    assert "UI baseline local report" in presentation.slides[0].shapes.title.text
    expect(cms_page.locator("#report-action-status")).to_have_attribute("data-state", "success")
    artifact = capture_page(cms_page, "workflow-report-success")
    download.save_as(artifact.with_suffix(".pptx"))

    cms_page.evaluate("""() => {
      globalThis.cmsCopied = [];
      Object.defineProperty(navigator, 'clipboard', {configurable: true, value: {
        write: async items => {
          for (const item of items) {
            globalThis.cmsCopied.push({
              html: await (await item.getType('text/html')).text(),
              text: await (await item.getType('text/plain')).text(),
            });
          }
        }
      }});
    }""")
    email.click()
    expect(cms_page.locator("#email-draft-dialog")).to_be_visible()
    EmailDraftResponse.model_validate(cms_side_effects.responses["email-draft"])
    expect(cms_page.locator("#email-draft-preview")).to_contain_text("UI baseline")
    expect(cms_page.locator("#email-draft-preview img")).to_have_count(2)
    cms_page.wait_for_function("() => [...document.querySelectorAll('#email-draft-preview img')].every(image => image.complete && image.naturalWidth > 0)")
    cms_page.locator("#email-draft-copy").click()
    expect(cms_page.locator("#email-draft-status")).to_contain_text("클립보드에 복사했습니다")
    copied = cms_page.evaluate("globalThis.cmsCopied")
    assert len(copied) == 1
    assert "UI baseline" in copied[0]["text"]
    assert 'data:image/png;base64,' in copied[0]["html"]
    calls = {item["endpoint"]: item for item in cms_side_effects.requests}
    for endpoint in ("diagnostic-ppt", "email-draft"):
        assert calls[endpoint]["request"]["recordId"] == real["fft"]["recordId"]
        assert calls[endpoint]["request"]["until"] == "2026-10-29"
    capture_page(cms_page, "workflow-email-success")


def test_local_report_error(cms_page, cms_side_effects):
    from playwright.sync_api import expect

    cms_side_effects.ppt_mode = "error"
    open_page(cms_page, "/dashboard2?dataset=1&until=2026-10-29")
    export = cms_page.locator("#diagnostic-ppt-export")
    expect(export).to_be_enabled()
    export.click()
    expect(cms_page.locator("#report-action-status")).to_have_attribute("data-state", "error")
    expect(cms_page.locator("#report-action-status")).to_contain_text("PPT generation disabled in UI baseline")
    expect(export).to_be_enabled()
    capture_page(cms_page, "workflow-report-error")


def test_local_chat_and_fail_closed_barrier(cms_page, cms_side_effects, cms_base_url):
    from playwright.sync_api import expect
    from app.chat import DashboardChatResponse

    with pytest.raises(AssertionError, match="writes/fetch are disabled"):
        cms_page.request.post("/api/dashboard/1/ai-analysis", data={})
    with pytest.raises(AssertionError, match="only permits known local"):
        cms_page.request.get("https://provider.invalid/api/paid")
    open_page(cms_page, "/dashboard2?dataset=1&until=2026-10-29")
    cms_page.locator("#chat-open").click()
    cms_page.locator("#chat-input").fill("Summarize the local fixture")
    cms_page.locator("#chat-send").click()
    expect(cms_page.locator("#chat-messages")).to_contain_text("UI baseline stub chat")
    DashboardChatResponse.model_validate(cms_side_effects.responses["chat"])
    capture_page(cms_page, "workflow-chat-success")
    other = cms_page.context.new_page()
    try:
        other.goto("/equipment-status2")
        cms_side_effects.expected_blocked.update({
          f"{cms_base_url}/api/not-registered",
          f"{cms_base_url}/api/not-registered-read",
        })
        results = other.evaluate("""async () => {
          const write = await fetch('/api/not-registered', {method:'POST'});
          const read = await fetch('/api/not-registered-read');
          const chat = await fetch('/api/dashboard/1/chat?barrier=1', {
            method:'POST', headers:{'Content-Type':'application/json'}, body:'{}'
          });
          let externalBlocked = false;
          try { await fetch('https://provider.invalid/api/paid', {method:'POST'}); }
          catch { externalBlocked = true; }
          return {write:write.status, read:read.status, chat:chat.status, externalBlocked};
        }""")
        assert results == {"write": 503, "read": 503, "chat": 200, "externalBlocked": True}
        assert len(cms_side_effects.blocked) == 3
    finally:
        other.close()