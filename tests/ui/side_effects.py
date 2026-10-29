from __future__ import annotations

import json
import re
from io import BytesIO
from urllib.parse import quote, urljoin, urlsplit


READ_PATH = re.compile(r"/api/(datasets|history|dashboard(?:/[^/]+(?:/fft/\d+|/related-history)?)?)")


class SafeSideEffects:
    def __init__(self, page, base_url):
        self.page = page
        self.base_url = base_url
        self.ppt_mode = "success"
        self.ai_mode = "success"
        self.ai_headline = None
        self.ai_summary = None
        self.ai_record_id = None
        self.pending_ai = []
        self.requests = []
        self.blocked = []
        self.expected_blocked = set()
        self.responses = {}
        self.guard_api_requests()
        page.context.route("**/*", self.handle)

    def guard_api_requests(self):
        api = self.page.context.request

        def guarded(method):
            def read(url, **kwargs):
                target = urlsplit(urljoin(self.base_url, url))
                base = urlsplit(self.base_url)
                if (target.scheme, target.netloc) != (base.scheme, base.netloc) or not READ_PATH.fullmatch(target.path):
                    raise AssertionError("UI APIRequestContext only permits known local non-paid reads")
                return method(url, **kwargs)
            return read

        def deny(*args, **kwargs):
            raise AssertionError("UI APIRequestContext writes/fetch are disabled; use browser-local fixtures")

        for name in ("get", "head"):
            setattr(api, name, guarded(getattr(api, name)))
        for name in ("post", "put", "patch", "delete", "fetch"):
            setattr(api, name, deny)

    def ai_response(self, dataset_id, request):
        params = {"until": request["until"]} if request.get("until") else {}
        response = self.page.context.request.get(f"/api/dashboard/{quote(dataset_id, safe='')}", params=params)
        assert response.ok, f"Non-paid dashboard fixture read failed: {response.status}"
        dashboard = response.json()
        assert dashboard["meta"]["datasetId"] == dataset_id
        record = dashboard.get("fft") or {}
        record_id = record.get("recordId")
        return {
            "status": "ready", "cached": True, "model": "stubbed-ui-baseline",
            "generatedAt": "2026-10-29T00:00:00+00:00", "fftSeverityAssessments": [],
            "analysis": {
                "headline": self.ai_headline or "UI baseline stub", "riskLevel": "unknown", "riskLabel": "Inspection required",
                "trendSummary": self.ai_summary or "UI baseline local trend evidence", "fftSummary": self.ai_summary or "UI baseline real FFT reference",
                "fftAnalysisPeriod": {"start": None, "end": None},
                "fftReferenceRecordIds": [self.ai_record_id or record_id] if record_id else [],
                "trendFindings": [{"heading": "Trend", "detail": "Local baseline evidence"}],
                "fftFindings": [{"heading": "FFT", "detail": "Selected dataset and cutoff reference"}],
                "historyFindings": [],
                "likelyCauses": [{"cause": "Unknown", "confidence": "low", "rationale": "Baseline only"}],
                "counterEvidence": ["No paid AI calls"],
                "recommendedActions": [{"priority": "planned", "action": "Inspect equipment", "rationale": "UI baseline"}],
            },
        }

    def release_ai(self):
        for route, response in self.pending_ai:
            route.fulfill(status=200, json=response)
        self.pending_ai.clear()
        self.ai_mode = "success"

    def local_pptx(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "UI baseline local report"
        slide.placeholders[1].text = "Deterministic UI fixture; no provider or validation service called."
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    def handle(self, route):
        request = route.request
        target = urlsplit(request.url)
        base = urlsplit(self.base_url)
        if (target.scheme, target.netloc) != (base.scheme, base.netloc):
            self.blocked.append({"url": request.url, "reason": "external request"})
            route.abort("blockedbyclient")
            return
        match = re.fullmatch(r"/api/dashboard/([^/]+)/(ai-analysis|chat|email-draft|diagnostic-ppt)", target.path)
        if match:
            dataset_id, endpoint = match.groups()
            payload = request.post_data_json or {}
            self.requests.append({"endpoint": endpoint, "dataset": dataset_id, "request": payload})
            if endpoint == "ai-analysis":
                if self.ai_mode == "error":
                    route.fulfill(status=503, json={"detail": {"message": "Task5 local AI unavailable"}})
                    return
                response = self.ai_response(dataset_id, payload)
                if self.ai_mode == "held":
                    self.responses[endpoint] = response
                    self.pending_ai.append((route, response))
                    return
            elif endpoint == "chat":
                response = {"message": {"role": "assistant", "content": "UI baseline stub chat"}, "model": "stubbed-ui-baseline", "toolName": None, "toolArguments": {}, "toolResult": {}, "toolCalls": []}
            elif endpoint == "email-draft":
                response = {"status": "ready", "model": "stubbed-ui-baseline", "draft": {
                    "subject": "UI baseline", "greeting": "Hello.", "status": "Local baseline validation.",
                    "trend": "Real server trend response.", "fft": "Real selected FFT response.",
                    "priorityActions": [{"priority": "planned", "action": "Review the baseline."}],
                    "limitations": ["No email is sent"], "closing": "Thank you.",
                }}
            elif self.ppt_mode == "error":
                route.fulfill(status=503, json={"detail": {"code": "provider_unavailable", "message": "PPT generation disabled in UI baseline"}})
                return
            else:
                route.fulfill(status=200, content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", headers={
                    "Content-Disposition": 'attachment; filename="ui-baseline-local.pptx"',
                    "X-PPT-Validation": quote(json.dumps({"status": "passed", "repairs": 0, "issues": []})),
                }, body=self.local_pptx())
                return
            self.responses[endpoint] = response
            route.fulfill(status=200, json=response)
        elif request.method not in {"GET", "HEAD"} or (target.path.startswith("/api/") and not READ_PATH.fullmatch(target.path)):
            self.blocked.append({"url": request.url, "reason": "unregistered API or write"})
            route.fulfill(status=503, json={"detail": "UI baseline blocked unregistered side effect"})
        else:
            route.continue_()