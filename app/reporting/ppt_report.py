"""Template-based one-slide CMS diagnostic PPTX generation."""

from __future__ import annotations

import math
import re
from collections.abc import Mapping
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

from app.reporting.ppt_ai import PptDiagnosis
from app.reporting.ppt_layout import PptLayout, PptRegion, default_ppt_layout


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "docs" / "ppt_result_sample.pptx"
FONT_NAME = "Malgun Gothic"
TITLE_COLOR = (15, 81, 127)
BODY_COLOR = (0, 0, 0)
ACTION_COLOR = (0, 0, 255)


def build_diagnostic_ppt(
    report_context: Mapping[str, Any],
    diagnosis: PptDiagnosis,
    trend_png: bytes | None,
    fft_png: bytes | None,
    layout: PptLayout | None = None,
) -> bytes:
    """Fill the supplied single-slide template with current diagnostic evidence."""
    if not TEMPLATE_PATH.is_file():
        raise ValueError("진단 보고서 PPT 템플릿을 찾을 수 없습니다.")

    presentation = Presentation(TEMPLATE_PATH)
    if not presentation.slides:
        raise ValueError("진단 보고서 PPT 템플릿에 슬라이드가 없습니다.")

    _remove_extra_slides(presentation)
    slide = presentation.slides[0]
    evidence = _evidence(report_context)
    equipment = _mapping(evidence.get("equipment"))
    trend = _mapping(evidence.get("trend"))
    fft = _mapping(evidence.get("fft"))
    metrics = _metrics_by_key(trend)

    tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
    if len(tables) < 2:
        raise ValueError("진단 보고서 PPT 템플릿의 표 구성을 확인할 수 없습니다.")
    _, metadata_table = tables[:2]
    layout = layout or default_ppt_layout()
    title = _find_title_shape(slide)
    trend_text, fft_text = _find_analysis_text_shapes(slide)
    picture_frames = _picture_frames(slide)
    if len(picture_frames) < 2:
        raise ValueError("진단 보고서 PPT 템플릿의 그래프 영역을 확인할 수 없습니다.")

    table_shapes = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
    table_shapes[0].name = "panel_background"
    table_shapes[1].name = "metadata"
    _apply_region(table_shapes[1], layout.metadata)
    panel = table_shapes[0]
    regions = (layout.trend_text, layout.fft_text, layout.trend_chart, layout.fft_chart)
    left = min(region.x for region in regions) - 0.1
    top = min(region.y for region in regions) - 0.4
    right = max(region.x + region.width for region in regions) + 0.1
    bottom = max(region.y + region.height for region in regions) + 0.1
    panel.left, panel.top = Inches(left), Inches(top)
    panel.width, panel.height = Inches(right - left), Inches(bottom - top)
    for column in panel.table.columns:
        column.width = Inches((right - left) / 2)
    panel.table.rows[0].height = Inches(0.32)
    panel.table.rows[1].height = Inches(bottom - top - 0.32)
    for row in panel.table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_NAME
                    run.font.size = Pt(12)
    for name, shape, region in (
        ("title", title, layout.title),
        ("trend_text", trend_text, layout.trend_text),
        ("fft_text", fft_text, layout.fft_text),
    ):
        shape.name = name
        _apply_region(shape, region)

    plant = _plain_text(equipment.get("plant"), 500)
    title_text = f"[{plant}] CMS Data 모니터링 및 진단결과" if plant else "CMS Data 모니터링 및 진단결과"
    _set_text(title, title_text, layout.title.font_size, TITLE_COLOR, bold=True)
    _fill_metadata_table(metadata_table, equipment, metrics, fft, diagnosis, layout.metadata.font_size)
    _set_text(trend_text, diagnosis.trend_finding, layout.trend_text.font_size, BODY_COLOR, bold=True)
    _set_fft_panel_text(fft_text, diagnosis, layout.fft_text.font_size)

    _clear_template_annotations(slide, (title, trend_text, fft_text))
    _replace_chart_images(slide, [
        tuple(Inches(value) for value in (region.x, region.y, region.width, region.height))
        for region in (layout.trend_chart, layout.fft_chart)
    ], trend_png, fft_png)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def diagnostic_ppt_filename_stem(report_context: Mapping[str, Any]) -> str:
    """Return an ASCII-safe filename stem for a generated diagnostic deck."""
    evidence = _evidence(report_context)
    equipment = _mapping(evidence.get("equipment"))
    fft = _mapping(evidence.get("fft"))
    equipment_id = _plain_text(equipment.get("id"), 48) or "equipment"
    safe_equipment_id = re.sub(r"[^A-Za-z0-9_-]+", "-", equipment_id).strip("-") or "equipment"
    measured_at = _plain_text(fft.get("measuredAt") or equipment.get("lastMeasuredAt"), 40)
    date_token = "".join(re.findall(r"\d", measured_at))[:8] or "report"
    return f"cms-diagnostic-{safe_equipment_id}-{date_token}"


def _remove_extra_slides(presentation: Presentation) -> None:
    while len(presentation.slides) > 1:
        slide_id = presentation.slides._sldIdLst[-1]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[-1]


def _evidence(report_context: Mapping[str, Any]) -> Mapping[str, Any]:
    evidence = report_context.get("evidence")
    return _mapping(evidence) if evidence is not None else report_context


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _plain_text(value: Any, limit: int) -> str:
    if value is None or isinstance(value, (Mapping, list, tuple, set)):
        return ""
    normalized = " ".join(str(value).split())
    return normalized[:limit] if normalized else ""


def _format_number(value: Any, digits: int = 2) -> str:
    if isinstance(value, bool):
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    if not math.isfinite(number):
        return ""
    if digits == 0:
        return f"{number:,.0f}"
    return f"{number:,.{digits}f}".rstrip("0").rstrip(".")


def _format_with_unit(value: Any, unit: str, digits: int = 2) -> str:
    formatted = _format_number(value, digits)
    return f"{formatted} {unit}" if formatted and unit else formatted


def _format_thresholds(metric: Mapping[str, Any]) -> str:
    warning = _format_number(metric.get("warningThreshold"))
    danger = _format_number(metric.get("dangerThreshold"))
    return f"{warning} / {danger}" if warning and danger else warning or danger


def _metrics_by_key(trend: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    metrics = trend.get("metrics")
    if not isinstance(metrics, list):
        return {}
    return {
        str(metric.get("key")): metric
        for metric in metrics
        if isinstance(metric, Mapping) and isinstance(metric.get("key"), str)
    }


def _find_title_shape(slide: Any) -> Any:
    candidates = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and not getattr(shape, "has_table", False)
    ]
    if not candidates:
        raise ValueError("진단 보고서 PPT 템플릿의 제목 영역을 확인할 수 없습니다.")
    return min(candidates, key=lambda shape: shape.top)


def _find_analysis_text_shapes(slide: Any) -> tuple[Any, Any]:
    candidates = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and not getattr(shape, "has_table", False)
        and shape.top > 2_000_000
    ]
    if len(candidates) < 2:
        raise ValueError("진단 보고서 PPT 템플릿의 분석 설명 영역을 확인할 수 없습니다.")
    candidates.sort(key=lambda shape: shape.left)
    return candidates[0], candidates[1]


def _picture_frames(slide: Any) -> list[tuple[int, int, int, int]]:
    frames = [
        (shape.left, shape.top, shape.width, shape.height)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    return sorted(frames)


def _fill_metadata_table(
    table: Any,
    equipment: Mapping[str, Any],
    metrics: Mapping[str, Mapping[str, Any]],
    fft: Mapping[str, Any],
    diagnosis: PptDiagnosis,
    font_size: float = 11,
) -> None:
    velocity = metrics.get("velocity", {})
    acceleration = metrics.get("accelerationAll", {})
    values = {
        (0, 1): _plain_text(equipment.get("plant"), 500),
        (0, 5): "",
        (0, 8): _format_report_timestamp(fft.get("measuredAt")),
        (0, 11): "",
        (1, 1): _plain_text(equipment.get("name") or equipment.get("displayName"), 500),
        (1, 8): "",
        (1, 11): _plain_text(equipment.get("sensorNumber") or equipment.get("sensor_number"), 500),
        (2, 2): _format_with_unit(equipment.get("rpm"), "RPM", 0),
        (2, 6): _plain_text(equipment.get("bearing"), 500),
        (2, 9): _format_thresholds(velocity),
        (2, 12): _format_thresholds(acceleration),
        (3, 1): _phenomenon_text(velocity, acceleration, diagnosis.trend_finding),
        (4, 1): _plain_text(diagnosis.headline, 90),
    }
    for (row, column), value in values.items():
        _set_table_cell_text(table.cell(row, column), value, font_size)


def _format_report_timestamp(value: Any) -> str:
    source = _plain_text(value, 40)
    if not source:
        return ""
    try:
        parsed = datetime.fromisoformat(source.replace("Z", "+00:00"))
    except ValueError:
        return source
    return f"’{parsed:%y}.{parsed.month}.{parsed.day} {parsed:%H:%M}"


def _phenomenon_text(
    velocity: Mapping[str, Any], acceleration: Mapping[str, Any], finding: str
) -> str:
    readings = [
        _format_with_unit(velocity.get("latestValue"), _plain_text(velocity.get("unit"), 12)),
        _format_with_unit(acceleration.get("latestValue"), _plain_text(acceleration.get("unit"), 12)),
    ]
    labels = ("속도", "가속도 ALL")
    current = ", ".join(
        f"{label} {reading}" for label, reading in zip(labels, readings, strict=True) if reading
    )
    return ". ".join(part for part in (current, _plain_text(finding, 280)) if part)


def _set_fft_panel_text(shape: Any, diagnosis: PptDiagnosis, size: float = 12) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    _add_paragraph(text_frame.paragraphs[0], diagnosis.fft_finding, size, BODY_COLOR, bold=True)
    action = next((item.action for item in diagnosis.recommended_actions if item.action), "")
    if action:
        _add_paragraph(text_frame.add_paragraph(), action, size, ACTION_COLOR)


def _set_text(shape: Any, value: str, size: float, color: tuple[int, int, int], bold: bool = False) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    _add_paragraph(text_frame.paragraphs[0], value, size, color, bold)


def _set_table_cell_text(cell: Any, value: str, font_size: float) -> None:
    text_frame = cell.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    _add_paragraph(text_frame.paragraphs[0], value, font_size, BODY_COLOR)


def _add_paragraph(paragraph: Any, value: str, size: float, color: tuple[int, int, int], bold: bool = False) -> None:
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_after = Pt(0)
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", "0")
    properties.set("indent", "0")
    for tag in ("buChar", "buAutoNum", "buNone"):
        for element in list(properties.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}")):
            properties.remove(element)
    properties.append(OxmlElement("a:buNone"))
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def _clear_template_annotations(slide: Any, retained_text_shapes: tuple[Any, ...]) -> None:
    """Remove sample images and frequency callouts after preserving text and tables."""
    retained_elements = {id(shape._element) for shape in retained_text_shapes}
    for shape in list(slide.shapes):
        if getattr(shape, "has_table", False) or id(shape._element) in retained_elements:
            continue
        shape._element.getparent().remove(shape._element)


def _replace_chart_images(
    slide: Any,
    frames: list[tuple[int, int, int, int]],
    trend_png: bytes | None,
    fft_png: bytes | None,
) -> None:
    for name, frame, image in zip(("trend_chart", "fft_chart"), frames[:2], (trend_png, fft_png), strict=True):
        if image:
            picture = slide.shapes.add_picture(_fit_png_to_frame(image, frame[2], frame[3]), *frame)
            picture.name = name


def _fit_png_to_frame(image: bytes, width: int, height: int) -> BytesIO:
    """Letterbox the complete graph so axes and legends are never cropped."""
    with Image.open(BytesIO(image)) as source:
        source = source.convert("RGB")
        target_width = 1600
        target_height = max(1, round(target_width * height / width))
        scale = min(target_width / source.width, target_height / source.height)
        source = source.resize((max(1, round(source.width * scale)), max(1, round(source.height * scale))), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (target_width, target_height), "white")
        canvas.paste(source, ((target_width - source.width) // 2, (target_height - source.height) // 2))
        output = BytesIO()
        canvas.save(output, format="PNG")
        output.seek(0)
        return output


def _apply_region(shape: Any, region: PptRegion) -> None:
    if getattr(shape, "has_table", False):
        table = shape.table
        old_width = sum(column.width for column in table.columns)
        for column in table.columns:
            column.width = round(column.width * Inches(region.width) / old_width)
        weights = (0.50, 0.55, 0.40, 0.40, 0.30)
        for row, weight in zip(table.rows, weights, strict=True):
            row.height = Inches(region.height * weight / sum(weights))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = FONT_NAME
                        run.font.size = Pt(region.font_size)
    shape.left, shape.top = Inches(region.x), Inches(region.y)
    shape.width, shape.height = Inches(region.width), Inches(region.height)
